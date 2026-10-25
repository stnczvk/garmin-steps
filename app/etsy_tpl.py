"""Szablony Etsy (PlanPerfect Sheets) budowane na serwerze – do personalizacji zamówień."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, ColorScaleRule, CellIsRule
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.comments import Comment
from openpyxl.worksheet.properties import PageSetupProperties
import datetime as dt

NAVY="1F2A44"; NAVY2="2C3A5C"; SAGE="5FA987"; SAGEL="E6F3EC"; CREAM="FAF8F4"; INPUT="FFF6DB"; LINE="E3E6EB"
INK="1F2937"; GREY="6B7280"; CORAL="E07A5F"; GOLD="E9B44C"; BLUE="5B8DEF"; WHITE="FFFFFF"
F=lambda **k: Font(name="Arial", **k)
fill=lambda c: PatternFill("solid", fgColor=c)
s=lambda c=LINE: Side(style="thin", color=c)
box=Border(bottom=s())
MONEY='#,##0.00;-#,##0.00;"–"'
N=500

DEF_INC=["Product sales","Services","Commissions","Subscriptions","Other income"]
DEF_EXP=["Materials & supplies","Shipping & packaging","Software & subscriptions","Marketing & ads","Platform fees","Equipment","Travel","Phone & internet","Rent / workspace","Professional services","Taxes & licences","Other"]
DEF_PAY=["Bank transfer","Card","Cash","PayPal","Stripe","Other"]


def build(PATH, NAME="Your Business Name", CURRENCY="USD", INC=None, EXP=None, PAY=None, YEAR=None, EXAMPLE=True):
    """Small Business Profit & Expense Tracker (PlanPerfect Sheets) – wersja ogólna albo spersonalizowana."""
    INC=(INC or DEF_INC)[:20]; EXP=(EXP or DEF_EXP)[:20]; PAY=(PAY or DEF_PAY)[:20]; YEAR=YEAR or dt.date.today().year
    wb=Workbook()

    def page(ws, cols, title, sub, crumb):
        ws.sheet_view.showGridLines=False
        ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True)
        ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
        ws.column_dimensions["A"].width=3
        for r in range(1,5):
            for c in range(1,cols+2): ws.cell(r,c).fill=fill(NAVY)
        for c in range(1,cols+2): ws.cell(5,c).fill=fill(SAGE)
        ws.row_dimensions[1].height=12; ws.row_dimensions[2].height=34; ws.row_dimensions[3].height=20; ws.row_dimensions[4].height=10; ws.row_dimensions[5].height=4
        ws["B2"]=title; ws["B2"].font=F(bold=True,size=22,color=WHITE); ws["B2"].alignment=Alignment(vertical="center")
        ws["B3"]=sub; ws["B3"].font=F(size=10,color="B8C4DA")
        c=ws.cell(2,cols,crumb); c.font=F(bold=True,size=9,color="B8C4DA"); c.alignment=Alignment(horizontal="right",vertical="center")
        ws.row_dimensions[6].height=14

    def card(ws, r, c, w, label, formula, color, fmt=MONEY, note=None):
        # pasek koloru, etykieta, duża liczba, podpis – karta na w kolumn
        for cc in range(c,c+w):
            ws.cell(r,cc).fill=fill(color)
            for rr in (r+1,r+2,r+3): ws.cell(rr,cc).fill=fill(CREAM)
        ws.row_dimensions[r].height=5
        ws.merge_cells(start_row=r+1,start_column=c,end_row=r+1,end_column=c+w-1)
        ws.merge_cells(start_row=r+2,start_column=c,end_row=r+2,end_column=c+w-1)
        ws.merge_cells(start_row=r+3,start_column=c,end_row=r+3,end_column=c+w-1)
        a=ws.cell(r+1,c,"  "+label.upper()); a.font=F(bold=True,size=8,color=GREY); a.alignment=Alignment(vertical="bottom")
        v=ws.cell(r+2,c,formula); v.font=F(bold=True,size=20,color=INK); v.number_format=fmt; v.alignment=Alignment(horizontal="left",vertical="center",indent=1)
        n=ws.cell(r+3,c,note or ""); n.font=F(size=8,color=GREY); n.alignment=Alignment(indent=1,vertical="top")
        ws.row_dimensions[r+1].height=20; ws.row_dimensions[r+2].height=32; ws.row_dimensions[r+3].height=18
        for rr in range(r,r+4):
            x=ws.cell(rr,c+w-1); x.border=Border(right=Side(style="thick",color=WHITE))

    def head(ws, r, c, text, w=1, color=NAVY):
        for cc in range(c,c+w): ws.cell(r,cc).fill=fill(color)
        x=ws.cell(r,c,text); x.font=F(bold=True,size=10,color=WHITE); x.alignment=Alignment(vertical="center",indent=1)
        ws.row_dimensions[r].height=24

    def section(ws, r, c, text):
        x=ws.cell(r,c,text); x.font=F(bold=True,size=12,color=NAVY)
        ws.row_dimensions[r].height=26; x.alignment=Alignment(vertical="bottom")

    # =================== START ===================
    st=wb.active; st.title="Start"
    page(st,8,"Small Business Profit & Expense Tracker","Know your real profit every month – in 5 minutes a week","PLANPERFECT SHEETS")
    st.column_dimensions["B"].width=6
    for col in "CDEFGH": st.column_dimensions[col].width=16
    section(st,8,2,"How to use")
    steps=[("1","Setup","Type your business name, currency, year and tax set-aside %. Rename the categories to fit your business."),
           ("2","Log income","Add every payment you receive on the Income tab. Mark it Paid or Pending."),
           ("3","Log expenses","Add every cost on the Expenses tab. Pick a category from the drop-down."),
           ("4","Check the Dashboard","Profit, margin, tax to set aside and where your money goes – all calculated for you."),
           ("5","Repeat weekly","5 minutes every Friday keeps your numbers accurate all year.")]
    for i,(n,t,d) in enumerate(steps):
        r=9+i*2
        b=st.cell(r,2,n); b.fill=fill(SAGE); b.font=F(bold=True,size=14,color=WHITE); b.alignment=Alignment(horizontal="center",vertical="center")
        st.merge_cells(start_row=r,start_column=3,end_row=r,end_column=8)
        x=st.cell(r,3,f"{t}  –  {d}"); x.font=F(size=11,color=INK); x.alignment=Alignment(vertical="center",indent=1,wrap_text=True)
        for cc in range(3,9): st.cell(r,cc).fill=fill(CREAM)
        st.row_dimensions[r].height=34; st.row_dimensions[r+1].height=6
    section(st,20,2,"Legend")
    st.cell(21,2).fill=fill(INPUT); st.cell(21,2).border=Border(left=s(GOLD),right=s(GOLD),top=s(GOLD),bottom=s(GOLD))
    st.cell(21,3,"Yellow cells = type here. Everything else is calculated for you.").font=F(size=10,color=INK)
    st.cell(22,2).fill=fill(CREAM); st.cell(22,3,"The first row on Income and Expenses is an example – overwrite it.").font=F(size=10,color=INK)
    section(st,24,2,"Good to know")
    for i,t in enumerate(["Pending invoices are shown separately and are not counted as income until marked Paid.",
                          "Each log holds 500 rows. Need a new year? File > Make a copy and change the year on Setup.",
                          "Tax set-aside is an estimate to help you save – check your local tax rules."]):
        st.cell(25+i,3,"•  "+t).font=F(size=10,color=GREY)

    # =================== SETUP ===================
    su=wb.create_sheet("Setup")
    page(su,9,"Setup","Your business details and categories","PLANPERFECT SHEETS")
    for col,w in zip("BCDEFGHI",[30,26,4,30,4,30,4,4]): su.column_dimensions[col].width=w
    section(su,7,2,"Business details")
    rows=[("Business name",NAME,None),("Currency (shown in headers)",CURRENCY,None),("Tracking year",YEAR,"0"),("Tax set-aside (% of profit)",0.25,"0%")]
    for i,(k,v,f) in enumerate(rows):
        r=8+i; su.row_dimensions[r].height=26
        a=su.cell(r,2,k); a.font=F(size=10,color=GREY); a.alignment=Alignment(vertical="center")
        c=su.cell(r,3,v); c.fill=fill(INPUT); c.font=F(bold=True,size=11,color=INK); c.alignment=Alignment(vertical="center",indent=1)
        c.border=Border(bottom=s(GOLD))
        if f: c.number_format=f
    su["C11"].comment=Comment("Rule of thumb for self-employed: 25–30% of profit. Check your local tax rules.","PlanPerfect Sheets")
    section(su,13,2,"Categories  (rename freely – drop-downs update automatically)")
    inc=INC
    exp=EXP
    pay=PAY
    for col,t,lst,colr in ((2,"Income categories",inc,SAGE),(5,"Expense categories",exp,CORAL),(7,"Payment methods",pay,BLUE)):
        head(su,14,col,t,1,colr)
        for i in range(20):
            c=su.cell(15+i,col,lst[i] if i<len(lst) else None); c.fill=fill(INPUT); c.font=F(size=10,color=INK); c.border=Border(bottom=s("F1E3B5"))
            c.alignment=Alignment(indent=1); su.row_dimensions[15+i].height=20

    # =================== LOGI ===================
    def log(name, sub, heads, widths, catrange, extra, example, accent):
        ws=wb.create_sheet(name); n=len(heads)
        page(ws,n+1,name,sub,"PLANPERFECT SHEETS")
        for i,w in enumerate(widths): ws.column_dimensions[chr(66+i)].width=w
        # pasek podsumowania
        Y="Setup!$C$10"; A=f"$F$12:$F${11+N}"; D=f"$B$12:$B${11+N}"
        ws.row_dimensions[7].height=18; ws.row_dimensions[8].height=26
        ws["B7"]="ENTRIES"; ws["D7"]="TOTAL THIS YEAR"; ws["F7"]="THIS MONTH"
        for a in ("B7","D7","F7"): ws[a].font=F(bold=True,size=8,color=GREY)
        ws["B8"]=f'=COUNT({A})'; ws["D8"]=f'=SUMIFS({A},{D},">="&DATE({Y},1,1),{D},"<"&DATE({Y}+1,1,1))'
        ws["F8"]=f'=SUMIFS({A},{D},">="&DATE(YEAR(TODAY()),MONTH(TODAY()),1),{D},"<"&DATE(YEAR(TODAY()),MONTH(TODAY())+1,1))'
        for a in ("B8","D8","F8"): ws[a].font=F(bold=True,size=16,color=accent); ws[a].number_format=MONEY if a!="B8" else "0"; ws[a].alignment=Alignment(horizontal="left")
        ws.row_dimensions[9].height=10
        for i,h in enumerate(heads):
            c=ws.cell(11,2+i,h); c.fill=fill(NAVY); c.font=F(bold=True,size=10,color=WHITE); c.alignment=Alignment(horizontal="center" if i!=1 and i!=3 else "left",vertical="center",indent=0 if i not in (1,3) else 1)
        ws["F11"]='="Amount ("&Setup!$C$9&")"'
        ws.row_dimensions[11].height=28
        for r in range(12,12+N):
            ws.row_dimensions[r].height=20
            for i in range(n):
                c=ws.cell(r,2+i); c.font=F(size=10,color=INK); c.border=Border(bottom=s("EEF0F3"))
                c.alignment=Alignment(horizontal="center" if i in (0,2,5,6) else ("right" if i==4 else "left"),vertical="center",indent=1 if i in (1,3) else 0)
            ws.cell(r,2).number_format="dd mmm yyyy"; ws.cell(r,6).number_format=MONEY
        for i,v in enumerate(example if EXAMPLE else []):
            c=ws.cell(12,2+i,v); c.fill=fill(INPUT)
        rng=f"B12:{chr(65+n)}{11+N}"
        ws.conditional_formatting.add(f"B13:{chr(65+n)}{11+N}",FormulaRule(formula=["MOD(ROW(),2)=1"],fill=fill(CREAM)))
        dv=DataValidation(type="list",formula1=catrange,allow_blank=True,showErrorMessage=False); ws.add_data_validation(dv); dv.add(f"D12:D{11+N}")
        for col,lst in extra:
            d=DataValidation(type="list",formula1=lst,allow_blank=True); ws.add_data_validation(d); d.add(f"{col}12:{col}{11+N}")
        dd=DataValidation(type="date",operator="greaterThan",formula1="DATE(2000,1,1)",allow_blank=True); ws.add_data_validation(dd); dd.add(f"B12:B{11+N}")
        ws.freeze_panes="C12"
        return ws
    inc_ws=log("Income","Every payment you receive – one row each",["Date","Client / source","Category","Description","Amount","Method","Status"],[14,26,20,34,15,16,13],
               "=Setup!$B$15:$B$34",[("G","=Setup!$G$15:$G$34"),("H",'"Paid,Pending"')],[dt.date(2026,1,8),"Example client","Services","Logo design project",450,"Bank transfer","Paid"],SAGE)
    inc_ws.conditional_formatting.add(f"H12:H{11+N}",CellIsRule(operator="equal",formula=['"Paid"'],font=F(bold=True,color=SAGE),fill=fill(SAGEL)))
    inc_ws.conditional_formatting.add(f"H12:H{11+N}",CellIsRule(operator="equal",formula=['"Pending"'],font=F(bold=True,color="B7791F"),fill=fill("FDF0D5")))
    exp_ws=log("Expenses","Every cost of running the business – one row each",["Date","Vendor","Category","Description","Amount","Deductible","Receipt"],[14,26,24,34,15,13,13],
               "=Setup!$E$15:$E$34",[("G",'"Yes,No"'),("H",'"Yes,No"')],[dt.date(2026,1,10),"Canva","Software & subscriptions","Monthly Pro plan",15,"Yes","Yes"],CORAL)

    # =================== DASHBOARD ===================
    db=wb.create_sheet("Dashboard",1)
    page(db,10,"Dashboard","Updates automatically from your Income and Expenses","PLANPERFECT SHEETS")
    db["B3"]='=Setup!$C$8&"   ·   "&Setup!$C$10&"   ·   amounts in "&Setup!$C$9'
    Y="Setup!$C$10"; IA=f"Income!$F$12:$F${11+N}"; ID=f"Income!$B$12:$B${11+N}"; IS=f"Income!$H$12:$H${11+N}"
    EA=f"Expenses!$F$12:$F${11+N}"; ED=f"Expenses!$B$12:$B${11+N}"; EC=f"Expenses!$D$12:$D${11+N}"
    yr=lambda D: f'{D},">="&DATE({Y},1,1),{D},"<"&DATE({Y}+1,1,1)'
    card(db,7,2,2,"Income (paid)",f'=SUMIFS({IA},{IS},"Paid",{yr(ID)})',SAGE,note="money actually received")
    card(db,7,4,2,"Expenses",f'=SUMIFS({EA},{yr(ED)})',CORAL,note="all business costs")
    card(db,7,6,3,"Net profit","=B9-D9",NAVY,note="income minus expenses")
    card(db,7,9,2,"Profit margin","=IF(B9=0,0,F9/B9)",BLUE,"0.0%","of every dollar earned")
    card(db,12,2,2,"Set aside for tax","=MAX(0,F9)*Setup!$C$11",GOLD,note='=TEXT(Setup!$C$11,"0%")&" of profit"')
    card(db,12,4,2,"Pending invoices",f'=SUMIFS({IA},{IS},"Pending")',GOLD,note=f'=COUNTIF({IS},"Pending")&" unpaid"')
    card(db,12,6,3,"Best month",'=IF(MAX(E21:E32)<=0,"–",INDEX(B21:B32,MATCH(MAX(E21:E32),E21:E32,0)))',SAGE,"@",'="profit "&TEXT(MAX(E21:E32),"#,##0")')
    card(db,12,9,2,"Avg monthly profit","=IFERROR(SUM(E21:E32)/MAX(1,SUMPRODUCT(--((C21:C32+D21:D32)>0))),0)",NAVY,note="per month with activity")
    section(db,18,2,"Month by month")
    for c,t in ((2,"Month"),(3,"Income"),(4,"Expenses"),(5,"Profit"),(6,"Margin")):
        pass
    hdr=[("B","Month"),("C","Income"),("D","Expenses"),("E","Profit"),("F","Margin")]
    # tabela miesięcy w kolumnach B..F (bez przerw)
    for col in "BCDEF": db.column_dimensions[col].width=14
    db.column_dimensions["G"].width=3
    for i,(col,t) in enumerate(hdr):
        c=db[f"{col}20"]; c.value=t; c.fill=fill(NAVY); c.font=F(bold=True,size=10,color=WHITE); c.alignment=Alignment(horizontal="left" if col=="B" else "right",vertical="center",indent=1)
    db.row_dimensions[20].height=24
    for m in range(1,13):
        r=20+m; db.row_dimensions[r].height=20
        db[f"B{r}"]=dt.date(2026,m,1).strftime("%B")
        mm=lambda D: f'{D},">="&DATE({Y},{m},1),{D},"<"&DATE({Y},{m}+1,1)'
        db[f"C{r}"]=f'=SUMIFS({IA},{IS},"Paid",{mm(ID)})'
        db[f"D{r}"]=f'=SUMIFS({EA},{mm(ED)})'
        db[f"E{r}"]=f"=C{r}-D{r}"; db[f"F{r}"]=f"=IF(C{r}=0,0,E{r}/C{r})"
        for col in "BCDEF":
            c=db[f"{col}{r}"]; c.font=F(size=10,color=INK); c.border=Border(bottom=s("EEF0F3"))
            c.number_format="0.0%" if col=="F" else MONEY; c.alignment=Alignment(horizontal="left" if col=="B" else "right",indent=1)
        if m%2==0:
            for col in "BCDEF": db[f"{col}{r}"].fill=fill(CREAM)
    r=33
    for col,f in (("B","Total"),("C","=SUM(C21:C32)"),("D","=SUM(D21:D32)"),("E","=SUM(E21:E32)"),("F","=IF(C33=0,0,E33/C33)")):
        c=db[f"{col}{r}"]; c.value=f; c.font=F(bold=True,size=10,color=NAVY); c.fill=fill(SAGEL); c.number_format="0.0%" if col=="F" else MONEY
        c.alignment=Alignment(horizontal="left" if col=="B" else "right",indent=1)
    db.row_dimensions[33].height=24
    db.conditional_formatting.add("E21:E32",CellIsRule(operator="lessThan",formula=["0"],font=F(bold=True,color=CORAL)))
    db.conditional_formatting.add("E21:E32",CellIsRule(operator="greaterThan",formula=["0"],font=F(bold=True,color=SAGE)))
    # kategorie wydatków (H..K)
    for col,w in zip("HIJK",[24,14,12,3]): db.column_dimensions[col].width=w
    section(db,18,8,"Where the money goes")
    for col,t in (("H","Expense category"),("I","Amount"),("J","Share")):
        c=db[f"{col}20"]; c.value=t; c.fill=fill(NAVY); c.font=F(bold=True,size=10,color=WHITE); c.alignment=Alignment(horizontal="left" if col=="H" else "right",vertical="center",indent=1)
    for i in range(12):
        r=21+i
        db[f"H{r}"]=f'=IF(Setup!$E${15+i}="","",Setup!$E${15+i})'
        db[f"I{r}"]=f'=IF(H{r}="","",SUMIFS({EA},{EC},H{r},{yr(ED)}))'
        db[f"J{r}"]=f'=IF(OR(H{r}="",$D$33=0),"",I{r}/$D$33)'
        for col in "HIJ":
            c=db[f"{col}{r}"]; c.font=F(size=10,color=INK); c.border=Border(bottom=s("EEF0F3")); c.alignment=Alignment(horizontal="left" if col=="H" else "right",indent=1)
        db[f"I{r}"].number_format=MONEY; db[f"J{r}"].number_format="0%"
    db.conditional_formatting.add("J21:J32",ColorScaleRule(start_type="num",start_value=0,start_color="FFFFFF",end_type="num",end_value=0.5,end_color="F4B9A8"))
    # średnia
    db["K12"]=None
    # wykresy
    ch=BarChart(); ch.type="col"; ch.title="Income vs expenses"; ch.height=7.5; ch.width=17; ch.gapWidth=60; ch.overlap=-10
    ch.add_data(Reference(db,min_col=3,min_row=20,max_row=32),titles_from_data=True); ch.add_data(Reference(db,min_col=4,min_row=20,max_row=32),titles_from_data=True)
    ch.set_categories(Reference(db,min_col=2,min_row=21,max_row=32)); ch.y_axis.majorGridlines=None; ch.legend.position="b"
    ch.series[0].graphicalProperties.solidFill=SAGE; ch.series[0].graphicalProperties.line.noFill=True
    ch.series[1].graphicalProperties.solidFill=CORAL; ch.series[1].graphicalProperties.line.noFill=True
    db.add_chart(ch,"B36")
    pie=PieChart(); pie.title="Expenses by category"; pie.height=7.5; pie.width=11
    pie.add_data(Reference(db,min_col=9,min_row=20,max_row=32),titles_from_data=True); pie.set_categories(Reference(db,min_col=8,min_row=21,max_row=32))
    pie.legend.position="r"
    db.add_chart(pie,"H36")
    for ws in wb.worksheets: ws.sheet_properties.tabColor={"Start":NAVY,"Dashboard":SAGE,"Income":SAGE,"Expenses":CORAL}.get(ws.title,"B8C4DA")
    wb.active=1
    wb.save(PATH)
    return PATH
